# -*- coding: utf-8 -*-
"""
Gera o documento de Requisitos do Sistema Shiva Zen (DOCX + PDF) a partir dos
dados extraidos do codigo real (docs/_dados_requisitos.json) e dos diagramas UML
renderizados em docs/diagramas/*.png.

Uso: python docs/gerar_requisitos.py
"""
import json
import os
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
DIAG = os.path.join(BASE, 'diagramas')
DADOS = os.path.join(BASE, '_dados_requisitos.json')

VERSAO = '2.3'
DATA = date(2026, 10, 28).strftime('%d/%m/%Y')
# (versao, data, descricao) — acrescente uma linha a cada regeneracao relevante
HISTORICO = [
    ('2.1', '21/06/2026', 'Documento de requisitos + modelagem UML gerado a partir do código '
                          '(pós-auditoria SWE de backend e migração do front-end para base_v2).'),
    ('2.2', '23/09/2026', 'Atualização pós-auditoria pré-produção: OTP obrigatório em todo agendamento, '
                          'consentimentos LGPD/art. 11 com prova de aceite imutável, 2FA obrigatório '
                          'para ADMIN, histórico do prontuário, agendamento interno, deploy por Docker '
                          '+ migração atômica (migrations 0040–0046).'),
    ('2.3', '28/10/2026', 'Atualização pós-revisão de outubro/2026 (migrations 0047–0061): termo de saúde, '
                          'agenda unificada (agenda_horario + agenda_excecao), links públicos em hash '
                          '(link_acesso), consentimentos com histórico, questionário do prontuário como '
                          'formulário PRONTUARIO, pacote congelado na venda (compra_pacote_item), comissão como '
                          'livro-razão, padrão SQL de nomes, lógica no banco (triggers, funções, views e '
                          'procedures), trilhas imutáveis, cifragem do dado de saúde, fila de e-mail na '
                          'notificação e saída do Celery (jobs.py/rodar_job); services/agenda e '
                          'services/atendimento, Ator da requisição, Django Forms no painel, wizard com botão '
                          'Continuar e PWA. Diagramas UML revisados e tabela de rotas.'),
]
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
DARK = RGBColor(0x4A, 0x34, 0x25)

with open(DADOS, encoding='utf-8') as f:
    D = json.load(f)

doc = Document()

# ---- estilos base ----
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)

for lvl, sz in ((1, 16), (2, 13)):
    st = doc.styles[f'Heading {lvl}']
    st.font.color.rgb = DARK
    st.font.size = Pt(sz)
    st.font.name = 'Calibri'


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(hdr[i], '7A5F1F')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(9)
    if widths:
        for col, w in enumerate(widths):
            for cell in t.columns[col].cells:
                cell.width = Inches(w)
    doc.add_paragraph()
    return t


def _orientar(paisagem):
    """Garante a orientacao da pagina atual. Secao nova (= pagina nova) so quando ela muda:
    voltar ao retrato logo apos cada figura deixava uma pagina em branco entre 2 paisagens."""
    sec = doc.sections[-1]
    if (sec.orientation == WD_ORIENT.LANDSCAPE) == paisagem:
        return False
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    return True


def add_image(nome, legenda, paisagem=False):
    path = os.path.join(DIAG, nome)
    if not os.path.exists(path):
        doc.add_paragraph(f'[diagrama ausente: {nome}]')
        return
    img = DocxImage.from_file(path)
    proporcao = img.px_height / img.px_width
    # Diagrama alto aproveita melhor a pagina em retrato (paisagem so p/ os largos).
    paisagem = paisagem and proporcao <= 0.8
    if paisagem:
        if not _orientar(True):
            doc.add_page_break()   # figura em paisagem logo apos outra
    else:
        _orientar(False)
    # Cabe na area util (largura e altura, com folga p/ a legenda): largura fixa
    # estourava a pagina com os diagramas altos.
    sec = doc.sections[-1]
    max_w = sec.page_width - sec.left_margin - sec.right_margin
    max_h = sec.page_height - sec.top_margin - sec.bottom_margin - Inches(0.8)
    largura = min(max_w, int(max_h / proporcao))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=largura)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(legenda)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def h1(txt):
    _orientar(False)
    doc.add_heading(txt, level=1)


def h2(txt):
    _orientar(False)
    doc.add_heading(txt, level=2)


def par(txt, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(txt)
    r.font.size = Pt(size)
    return p


# ===================== CAPA =====================
for _ in range(4):
    doc.add_paragraph()
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run('SHIVA ZEN')
r.bold = True
r.font.size = Pt(34)
r.font.color.rgb = GOLD
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Sistema de Gestão para Clínica de Estética')
r.font.size = Pt(15)
r.font.color.rgb = DARK
sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run('Jaqueline Aranha Estética')
r.font.size = Pt(12)
r.italic = True
for _ in range(2):
    doc.add_paragraph()
tit = doc.add_paragraph()
tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tit.add_run('Documento de Especificação de Requisitos\ne Modelagem UML')
r.bold = True
r.font.size = Pt(16)
for _ in range(6):
    doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(f'Versão {VERSAO}  |  {DATA}').font.size = Pt(11)

# ===================== SUMARIO =====================
doc.add_page_break()
h1('Sumário')
toc_p = doc.add_paragraph()
run = toc_p.add_run()
fldChar = OxmlElement('w:fldChar'); fldChar.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
instr.text = 'TOC \\o "1-2" \\h \\z \\u'
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
t_run = OxmlElement('w:t'); t_run.text = 'Atualize este campo no Word (Ctrl+A, F9) para gerar o sumário.'
fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
for el in (fldChar, instr, fldChar2, t_run, fldChar3):
    run._r.append(el)

# ===================== 1. INTRODUCAO =====================
doc.add_page_break()
h1('1. Introdução')
h2('1.1 Objetivo')
par('Este documento especifica os requisitos funcionais e não funcionais, as regras '
    'de negócio e a modelagem UML do sistema Shiva Zen — plataforma web de gestão da '
    'clínica de estética Jaqueline Aranha. O conteúdo foi derivado diretamente do '
    'código-fonte da aplicação (Django 5.2, migrations até 0061), garantindo aderência ao '
    'comportamento realmente implementado.')
h2('1.2 Escopo')
par('O sistema cobre o agendamento online público (sem senha: a cliente é identificada '
    'pelo celular, confirmado por código OTP via SMS em todo agendamento), com aceite da '
    'Política de Privacidade, consentimento para dado de saúde, termos do procedimento e '
    'autorizações de comunicação com histórico; a área "Meus agendamentos" (reagendar e '
    'cancelar); o portal do profissional; o painel administrativo (agenda com grade e '
    'exceções, agendamento interno, clientes, prontuário com histórico, pacotes, promoções, '
    'comissões em livro-razão, NPS e depoimentos, termos, auditoria); rotinas automáticas '
    '(fila de e-mail, lembrete D-1, NPS, pacotes, aniversário, limpeza de status, retenção '
    'LGPD) e integrações externas (SMS Zenvia, WhatsApp Meta, e-mail, Google Calendar, '
    'Cloudflare Turnstile, Sentry).')
h2('1.3 Visão geral da arquitetura')
par('Aplicação Django monolítica renderizada no servidor e servida por gunicorn '
    '(1 worker × 4 threads) em contêiner Docker no Railway, com PostgreSQL 18 como banco '
    'principal: 36 tabelas no padrão SQL, invariantes garantidas por EXCLUDE, CHECK, UNIQUE '
    'parciais e triggers, trilha de alterações por trigger, funções e views de relatório e '
    'procedures que concentram a retenção LGPD; o dado de saúde é gravado cifrado. Não há '
    'Celery: os jobs periódicos são funções disparadas por um cron externo via HTTP '
    'autenticado e o e-mail transacional é uma fila na própria tabela de notificações; cache '
    'e sessão são locais, com Redis opcional. O front-end usa Vite, Tailwind CSS v4, '
    'Alpine.js (build CSP) e componentes django-cotton, com PWA. A regra de negócio reside '
    'nos serviços — services/agenda concentra as regras de horário e services/atendimento os '
    'efeitos de cada mudança de status, aplicados na mesma transação da máquina de estados '
    '(FSM) do atendimento. Detalhes nas seções 8 e 10.')

# ===================== 2. ATORES =====================
doc.add_page_break()
h1('2. Atores do Sistema')
par('Atores identificados a partir dos papéis de acesso e dos fluxos públicos e internos.')
add_table(
    ['Ator', 'Descrição'],
    [[a['nome'], a['descricao']] for a in D['casos-uso']['atores']],
    widths=[1.7, 5.0],
)

# ===================== 3. REQUISITOS FUNCIONAIS =====================
doc.add_page_break()
h1('3. Requisitos Funcionais')
add_table(
    ['ID', 'Ator', 'Descrição'],
    [[rf['id'], rf['ator'], rf['descricao']] for rf in D['regras-requisitos']['requisitos_funcionais']],
    widths=[0.6, 1.4, 4.7],
)

# ===================== 4. REQUISITOS NAO-FUNCIONAIS =====================
doc.add_page_break()
h1('4. Requisitos Não Funcionais')
add_table(
    ['ID', 'Categoria', 'Descrição'],
    [[r['id'], r['categoria'], r['descricao']] for r in D['regras-requisitos']['requisitos_nao_funcionais']],
    widths=[0.6, 1.2, 4.9],
)

# ===================== 5. REGRAS DE NEGOCIO =====================
doc.add_page_break()
h1('5. Regras de Negócio')
add_table(
    ['ID', 'Título', 'Descrição'],
    [[r['id'], r['titulo'], r['descricao']] for r in D['regras-requisitos']['regras']],
    widths=[0.6, 1.6, 4.5],
)

# ===================== 6. CASOS DE USO =====================
doc.add_page_break()
h1('6. Casos de Uso')
add_image('01-casos-uso-publico.png', 'Figura 1 — Diagrama de Casos de Uso (UML): Área Pública (Cliente).')
add_image('02-casos-uso-interno.png', 'Figura 2 — Diagrama de Casos de Uso (UML): Área Interna (ADMIN, Profissional, Sistema).')
h2('6.1 Especificação dos casos de uso')
add_table(
    ['ID', 'Caso de Uso', 'Atores', 'Fluxo principal'],
    [[uc['id'], uc['nome'], ', '.join(uc['atores']), uc['fluxo']] for uc in D['casos-uso']['casos_de_uso']],
    widths=[0.5, 1.4, 1.5, 3.3],
)

# ===================== 7. MODELO DE DOMINIO =====================
doc.add_page_break()
h1('7. Modelo de Domínio (Diagrama de Classes)')
par('Os atributos seguem os nomes das classes Python; no banco, tabelas e colunas seguem o '
    'padrão SQL do projeto (por exemplo, ativo → eh_ativo, buffer_minutos → '
    'intervalo_minutos). O nome físico de cada tabela está no dicionário de entidades (7.1).')
add_image('03-classe-agendamento.png', 'Figura 3 — Diagrama de Classes (UML): Núcleo de Agendamento.', paisagem=True)
add_image('04-classe-financeiro.png', 'Figura 4 — Diagrama de Classes (UML): Pacotes, Carteira, Fidelidade e Comissão.', paisagem=True)
add_image('05-classe-prontuario-lgpd.png', 'Figura 5 — Diagrama de Classes (UML): Prontuário, Formulários, NPS, Termos, '
          'Consentimentos, Acesso e Trilhas.', paisagem=True)
h2('7.1 Dicionário de entidades')
linhas = []
for grupo in ('models-core', 'models-aux'):
    for m in D[grupo]['models']:
        linhas.append([m['nome'], m.get('tabela', ''), m['descricao']])
add_table(['Entidade', 'Tabela', 'Descrição'], linhas, widths=[1.5, 1.4, 3.8])

# ===================== 8. DIAGRAMA DE ESTADOS =====================
doc.add_page_break()
h1('8. Diagrama de Estados — Atendimento (FSM)')
par('O status do atendimento é governado por uma máquina de estados finitos; transições '
    'inválidas levantam exceção tipada (Atendimento.TransicaoInvalida). Estado inicial: '
    f"{D['fsm']['estado_inicial']}.")
add_image('06-estados-atendimento.png', 'Figura 6 — Diagrama de Estados (UML): ciclo de vida do Atendimento.', paisagem=True)
h2('8.1 Tabela de transições')
trans = [[t['de'], t['para'], t['acao']] for t in D['fsm']['transicoes'] if t['para'] != '(nenhum)']
add_table(['Estado origem', 'Estado destino', 'Ação / gatilho'], trans, widths=[1.6, 1.6, 3.5])
h2('8.2 Efeitos automáticos (services/atendimento / jobs)')
for ef in D['fsm']['efeitos']:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(ef).font.size = Pt(9)

# ===================== 9. DIAGRAMA DE SEQUENCIA =====================
doc.add_page_break()
h1('9. Diagrama de Sequência — Agendamento com OTP')
par('Fluxo do agendamento público: verificação do celular por OTP via SMS (exigida em todo '
    'agendamento), aceites e consentimentos com prova, controle de concorrência do horário e '
    'criação transacional do atendimento PENDENTE.')
add_image('07-sequencia-agendamento-otp.png', 'Figura 7 — Diagrama de Sequência (UML): agendamento público com OTP.', paisagem=True)
h2('9.1 Passo a passo')
for passo in D['arquitetura']['fluxo_agendamento_otp']:
    p = doc.add_paragraph(style='List Number')
    p.add_run(passo).font.size = Pt(9)

# ===================== 10. ARQUITETURA =====================
doc.add_page_break()
h1('10. Arquitetura e Integrações')
add_image('08-componentes-arquitetura.png', 'Figura 8 — Diagrama de Componentes / Implantação (UML).', paisagem=True)
h2('10.1 Componentes')
add_table(['Componente', 'Papel'],
          [[c['nome'], c['papel']] for c in D['arquitetura']['componentes']],
          widths=[1.8, 4.9])
h2('10.2 Integrações externas')
add_table(['Serviço', 'Uso'],
          [[i['nome'], i['uso']] for i in D['arquitetura']['integracoes_externas']],
          widths=[1.8, 4.9])
h2('10.3 Rotas')
add_table(['Área', 'Rota', 'Função'],
          [[r['grupo'], r['rota'], r['funcao']] for r in D['arquitetura']['rotas']],
          widths=[1.0, 2.9, 2.8])

# ===================== HISTORICO =====================
doc.add_page_break()
h1('Histórico de Versões')
add_table(['Versão', 'Data', 'Descrição'],
          [list(linha) for linha in HISTORICO],
          widths=[0.9, 1.2, 4.6])

# ---- rodape com numero de pagina ----
# As secoes criadas p/ as paginas em paisagem herdam (linked) o rodape da 1a:
# escrever em todas repetia o texto uma vez por secao no mesmo rodape.
for i, section in enumerate(doc.sections):
    footer = section.footer
    if i > 0 and footer.is_linked_to_previous:
        continue
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run('Shiva Zen — Requisitos v' + VERSAO + '   |   Pág. ')
    run.font.size = Pt(8)
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

out_docx = os.path.join(BASE, 'Requisitos-ShivaZen.docx')
doc.save(out_docx)
print('DOCX:', out_docx)

# ---- PDF ----
try:
    from docx2pdf import convert
    out_pdf = os.path.join(BASE, 'Requisitos-ShivaZen.pdf')
    convert(out_docx, out_pdf)
    print('PDF :', out_pdf)
except Exception as e:
    print('PDF falhou (gere abrindo o DOCX no Word > Salvar como PDF):', repr(e))
