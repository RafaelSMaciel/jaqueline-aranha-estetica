# -*- coding: utf-8 -*-
"""Planilha de custos operacionais: clínica com até 60 atendimentos por mês, ferramentas recomendadas.

Uma definição só (ABAS, abaixo) gera:
  - docs/tcc/ShivaZen-Custos-Operacionais.xlsx (openpyxl; no Windows o Excel recalcula e grava
    os valores, para quem abre a planilha sem recálculo ver os números);
  - com --json <pasta>: os valores (com fórmulas), a formatação e as mesclagens de cada aba em
    JSON, para montar a mesma planilha no Google Sheets.

Toda conta é fórmula sobre as premissas do topo da aba Custos (entradas em azul). Os preços e as
premissas são os de docs/CUSTOS-OPERACIONAIS.md: mudou um, mude os dois.

Uso (python do sistema, com openpyxl):
    python docs/tcc/gerar_planilha_custos.py [--json <pasta>] [--separador ";"]
"""
import argparse
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TCC = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(TCC, 'ShivaZen-Custos-Operacionais.xlsx')

FONTE = 'Arial'
CAFE = '#1F1A14'          # cromo do painel (DESIGN-E-MARKETING.md)
AREIA = '#F1EFE8'
AZUL = '#0000FF'          # entrada
CINZA = '#BFBFBF'
REAIS = '"R$" #,##0.00;("R$" #,##0.00);"-"'
REAIS_3 = '"R$" #,##0.000;("R$" #,##0.000);"-"'
DOLARES = '"US$" #,##0.00;("US$" #,##0.00);"-"'
PORCENTO = '0.0%;(0.0%);"-"'
INTEIRO = '#,##0'
DECIMAL = '#,##0.00'
CONSULTA = '30/10/2026'


def cabecalho(faixa):
    return {'range': faixa, 'bold': True, 'fg': '#FFFFFF', 'bg': CAFE, 'valign': 'middle', 'wrap': True}


def titulo(celula):
    return {'range': celula, 'bold': True, 'size': 14}


# ─── Custos ─────────────────────────────────────────────────────────────────────────
# Linhas 4–16: premissas (B = valor; observação mesclada em C:F). Linhas 18–27: custo mensal.
# Linhas 29–33: implantação. Linhas 35–37: o que ainda depende de mudar o código.
CUSTOS = [
    ['Custo do sistema: clínica com até 60 atendimentos por mês'],
    ['Células em azul são entradas: troque o valor e a planilha recalcula. Preços de lista consultados '
     f'em {CONSULTA}; endereços na aba Fontes.'],
    [],
    ['Premissa', 'Valor', 'Observação'],
    ['Atendimentos por mês', 60, 'Teto informado pela clínica'],
    ['Códigos de acesso por atendimento', 1, 'Um código por pedido online ou por acesso à área da cliente'],
    ['Mensagens de WhatsApp por atendimento', 1.25,
     'Lembrete D-1, pesquisa NPS (no máximo a cada 90 dias por cliente) e aviso de vaga'],
    ['Dólar (R$)', 5.10, 'PTAX de 16/10/2026 (R$ 5,08), arredondada para cima'],
    ['IOF na compra internacional no cartão', 0.035, 'Alíquota em vigor desde 2025'],
    ['Dólar com IOF (R$)', '=B8*(1+B9)', 'Usado nas contas em dólar'],
    ['Railway Hobby, mensalidade (US$)', 5, 'A mensalidade já inclui US$ 5 de uso'],
    ['Railway, uso estimado (US$/mês)', 7, 'Produção ~US$ 6 + ambiente de testes dormindo ~US$ 1'],
    ['Domínio .com.br (R$/ano)', 40, 'Registro.br'],
    ['SMS pré-pago (R$ por SMS)', 0.12, 'SMSGo, menor pacote; créditos não expiram'],
    ['WhatsApp, mensagem de utilidade (R$)', 0.035, 'Tarifa da Meta para o Brasil desde 01/10/2026'],
    ['Zenvia, pacote mínimo de SMS (R$/mês)', 100, 'Provedor que o código usa hoje (até 1.000 SMS)'],
    [],
    ['Custo mensal', 'Ferramenta', 'Quantidade por mês', 'Preço unitário (R$)', 'Custo mensal (R$)', 'Observação'],
    ['Hospedagem do site, do painel, do banco e das rotinas', 'Railway Hobby', 1, '=MAX(B11,B12)*B10',
     '=C19*D19', 'Inclui o PostgreSQL e o ambiente de testes'],
    ['Domínio e DNS', 'Registro.br', 1, '=B13/12', '=C20*D20', 'Pago uma vez por ano; aqui dividido por 12'],
    ['Código de acesso da cliente', 'SMS pré-pago (SMSGo)', '=B5*B6', '=B14', '=C21*D21',
     'Depende de trocar a Zenvia no código (ver "No radar")'],
    ['Lembrete, pesquisa NPS e aviso de vaga', 'WhatsApp Cloud API (Meta)', '=B5*B7', '=B15', '=C22*D22',
     'Só para quem autorizou o WhatsApp'],
    ['E-mail (transacional e campanhas)', 'Brevo Free', 1, 0, '=C23*D23',
     'Até 300 e-mails por dia; escolher na implantação (só configuração)'],
    ['Backup, erros, monitor das rotinas, antirrobô e avisos à equipe',
     'Cloudflare R2, Sentry, healthchecks.io, Turnstile e Web Push', 1, 0, '=C24*D24', 'Planos gratuitos'],
    ['Total por mês', None, None, None, '=SUM(E19:E24)', None],
    ['Por atendimento', None, None, None, '=IF(B5=0,0,E25/B5)', None],
    ['Por ano', None, None, None, '=E25*12', None],
    [],
    ['Implantação (uma vez)', 'Fornecedor', 'Quantidade', 'Preço unitário (R$)', 'Valor (R$)', 'Observação'],
    ['Domínio .com.br, 1º ano', 'Registro.br', 1, '=B13', '=C30*D30', 'Nos anos seguintes entra no custo mensal'],
    ['Primeiro crédito de SMS', 'SMSGo', 300, '=B14', '=C31*D31', 'Créditos que não expiram'],
    ['Contas nos serviços', 'Railway, Brevo, Meta, Cloudflare, Sentry, healthchecks.io', 1, 0, '=C32*D32',
     'Sem custo; a Meta exige CNPJ e o Railway pede cartão'],
    ['Total da implantação', None, None, None, '=SUM(E30:E32)', None],
    [],
    ['No radar: mudança de código', 'Hoje', 'Recomendado', 'Situação', 'Economia por mês (R$)', 'Observação'],
    ['Código de acesso por SMS', 'Zenvia, pacote mínimo mensal', 'SMS pré-pago', 'Pendente', '=B16-E21',
     'Trocar o provedor no módulo utils/sms; o custo mensal acima já considera a troca'],
    ['E-mail', 'Sem provedor configurado', 'Brevo Free', 'Na implantação', None,
     'Só configuração: EMAIL_BACKEND do Brevo e a chave BREVO_API_KEY'],
]
FMT_CUSTOS = {
    'rules': [
        titulo('A1'), {'range': 'A2', 'italic': True},
        cabecalho('A4:F4'), cabecalho('A18:F18'), cabecalho('A29:F29'), cabecalho('A35:F35'),
        {'range': 'B5:B7', 'fg': AZUL, 'number': DECIMAL}, {'range': 'B5', 'number': INTEIRO},
        {'range': 'B8', 'fg': AZUL, 'number': REAIS},
        {'range': 'B9', 'fg': AZUL, 'number': PORCENTO},
        {'range': 'B10', 'number': '"R$" #,##0.0000'},
        {'range': 'B11:B12', 'fg': AZUL, 'number': DOLARES},
        {'range': 'B13', 'fg': AZUL, 'number': REAIS},
        {'range': 'B14:B15', 'fg': AZUL, 'number': REAIS_3},
        {'range': 'B16', 'fg': AZUL, 'number': REAIS},
        {'range': 'C19:C24', 'number': INTEIRO},
        {'range': 'D19:E27', 'number': REAIS}, {'range': 'D21:D22', 'number': REAIS_3},
        {'range': 'D23:D24', 'fg': AZUL},
        {'range': 'A25:F25', 'bold': True, 'bg': AREIA}, {'range': 'A26:F27', 'bold': True},
        {'range': 'C30:C32', 'number': INTEIRO}, {'range': 'C31', 'fg': AZUL},
        {'range': 'D30:E33', 'number': REAIS}, {'range': 'D31', 'number': REAIS_3},
        {'range': 'D32', 'fg': AZUL},
        {'range': 'A33:F33', 'bold': True, 'bg': AREIA},
        {'range': 'E36', 'number': REAIS, 'bold': True},
        {'range': 'A19:F37', 'wrap': True, 'valign': 'top'},
        {'range': 'A4:F16', 'borders': CINZA}, {'range': 'A18:F27', 'borders': CINZA},
        {'range': 'A29:F33', 'borders': CINZA}, {'range': 'A35:F37', 'borders': CINZA},
    ],
    'merges': [f'C{lin}:F{lin}' for lin in range(4, 17)],
    'col_widths': {'A': 330, 'B': 250, 'C': 120, 'D': 140, 'E': 140, 'F': 380},
}

# ─── Fontes ─────────────────────────────────────────────────────────────────────────
FONTES = [
    ['Fontes'],
    [],
    ['Item', 'Fonte', 'Endereço', 'Consultado em'],
    ['Railway: planos e tarifas', 'Railway', 'https://railway.com/pricing', CONSULTA],
    ['Railway: modo Serverless (ambiente de testes)', 'Railway Docs', 'https://docs.railway.com/reference/app-sleeping', CONSULTA],
    ['Domínio .com.br', 'Terceiro', 'https://www.nuvemshop.com.br/blog/registro-br', CONSULTA],
    ['SMS pré-pago sem mensalidade', 'SMSGo', 'https://smsgo.com.br/', CONSULTA],
    ['SMS para OTP no Brasil: comparativo', 'Terceiro', 'https://x-apps.com.br/sms-verificacao-otp-brasil-provedores-precos/', CONSULTA],
    ['Zenvia: pacotes de SMS', 'Zenvia', 'https://zenvia.com/zoe/', CONSULTA],
    ['WhatsApp: modelo de cobrança', 'Meta for Developers', 'https://developers.facebook.com/docs/whatsapp/pricing', CONSULTA],
    ['WhatsApp: tarifas do Brasil (01/10/2026)', 'Tabela da Meta reproduzida por terceiro', 'https://zynfo.ai/whatsapp-pricing/brazil', CONSULTA],
    ['Brevo: plano gratuito (300 e-mails/dia)', 'Resenha de terceiro', 'https://www.sender.net/reviews/brevo/pricing/free-plan', CONSULTA],
    ['Cloudflare R2: cota gratuita', 'Cloudflare Docs', 'https://developers.cloudflare.com/r2/pricing', CONSULTA],
    ['Sentry: plano Developer', 'Terceiro', 'https://costbench.com/software/developer-tools/sentry/free-plan/', CONSULTA],
    ['healthchecks.io: planos', 'healthchecks.io', 'https://healthchecks.io/pricing/', CONSULTA],
    ['Câmbio: PTAX de 16/10/2026', 'Terceiro', 'https://www.simianx.ai/pt/stories/us-dollar-exchange-rate-by-year-every-year-1971-2026-table', CONSULTA],
    ['Câmbio: cotação oficial', 'Banco Central', 'https://www.bcb.gov.br/estabilidadefinanceira/historicocotacoes', CONSULTA],
]
FMT_FONTES = {
    'rules': [titulo('A1'), cabecalho('A3:D3'), {'range': 'A3:D17', 'borders': CINZA}],
    'col_widths': {'A': 320, 'B': 250, 'C': 520, 'D': 110},
}

ABAS = [
    ('Custos', CUSTOS, FMT_CUSTOS),
    ('Fontes', FONTES, FMT_FONTES),
]


# ─── xlsx ───────────────────────────────────────────────────────────────────────────
def _col(letras):
    n = 0
    for ch in letras:
        n = n * 26 + ord(ch) - 64
    return n


def _celulas(faixa):
    """'B5:D7' -> [(linha, coluna), ...] (1-based)."""
    m = re.match(r'^([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$', faixa)
    c1, l1, c2, l2 = m.group(1), int(m.group(2)), m.group(3) or m.group(1), int(m.group(4) or m.group(2))
    return [(lin, c) for lin in range(l1, l2 + 1) for c in range(_col(c1), _col(c2) + 1)]


def _hex(cor):
    return cor.lstrip('#').upper()


def _aplicar(ws, fmt):
    for faixa in fmt.get('merges', []):
        ws.merge_cells(faixa)
    for regra in fmt['rules']:
        for lin, col in _celulas(regra['range']):
            cel = ws.cell(row=lin, column=col)
            f = cel.font
            cel.font = Font(name=FONTE, size=regra.get('size', f.size or 10),
                            bold=regra.get('bold', f.bold), italic=regra.get('italic', f.italic),
                            color=_hex(regra['fg']) if 'fg' in regra else f.color)
            if 'bg' in regra:
                cel.fill = PatternFill('solid', fgColor=_hex(regra['bg']))
            if 'number' in regra:
                cel.number_format = regra['number']
            if any(k in regra for k in ('align', 'valign', 'wrap')):
                a = cel.alignment
                cel.alignment = Alignment(horizontal=regra.get('align', a.horizontal),
                                          vertical={'middle': 'center'}.get(regra.get('valign'), regra.get('valign', a.vertical)),
                                          wrap_text=regra.get('wrap', a.wrap_text))
            if 'borders' in regra:
                lado = Side(style='thin', color=_hex(regra['borders']))
                cel.border = Border(left=lado, right=lado, top=lado, bottom=lado)
    for faixa, px in fmt.get('col_widths', {}).items():
        a, _, b = faixa.partition(':')
        for n in range(_col(a), _col(b or a) + 1):
            ws.column_dimensions[get_column_letter(n)].width = round(px / 7, 1)


def gerar_xlsx():
    wb = Workbook()
    wb.remove(wb.active)
    for nome, linhas, fmt in ABAS:
        ws = wb.create_sheet(nome)
        for i, linha in enumerate(linhas, start=1):
            for j, valor in enumerate(linha, start=1):
                if valor is not None:
                    ws.cell(row=i, column=j, value=valor).font = Font(name=FONTE, size=10)
        _aplicar(ws, fmt)
        ws.sheet_view.showGridLines = False
    wb.save(SAIDA)
    return SAIDA


def recalcular_no_excel(caminho):
    """Abre no Excel (COM), recalcula e salva: o arquivo passa a ter os valores das fórmulas."""
    try:
        import win32com.client
    except ImportError:
        print('Excel indisponivel: o arquivo sai sem valores calculados (recalcula ao abrir)')
        return
    excel = win32com.client.DispatchEx('Excel.Application')
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(os.path.abspath(caminho))
        excel.CalculateFull()
        for ws in wb.Worksheets:
            ws.UsedRange.Rows.AutoFit()   # altura das linhas com texto quebrado
        wb.Save()
        wb.Close(False)
    finally:
        excel.Quit()


# ─── Google Sheets ──────────────────────────────────────────────────────────────────
def _para_sheets(valor, separador):
    if isinstance(valor, str):
        if valor.startswith('='):
            return valor.replace(',', separador) if separador != ',' else valor
        if re.match(r'^\d{2}/\d{2}/\d{4}$', valor):
            return "'" + valor          # data como texto
    return valor


def exportar_json(pasta, separador):
    os.makedirs(pasta, exist_ok=True)
    for i, (nome, linhas, fmt) in enumerate(ABAS):
        largura = max(len(linha) for linha in linhas)
        grade = [[_para_sheets(v, separador) for v in linha] + [None] * (largura - len(linha)) for linha in linhas]
        with open(os.path.join(pasta, f'{i}-valores.json'), 'w', encoding='utf-8') as f:
            json.dump({'aba': nome, 'valores': grade}, f, ensure_ascii=False)
        spec = {'rules': [dict(r, font=FONTE) for r in fmt['rules']], 'col_widths': fmt.get('col_widths', {})}
        with open(os.path.join(pasta, f'{i}-formato.json'), 'w', encoding='utf-8') as f:
            json.dump(spec, f, ensure_ascii=False)
        with open(os.path.join(pasta, f'{i}-mesclar.json'), 'w', encoding='utf-8') as f:
            json.dump(fmt.get('merges', []), f)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--json', help='pasta para os JSON do Google Sheets')
    ap.add_argument('--separador', default=',', help='separador de argumentos das fórmulas no Sheets')
    args = ap.parse_args()
    caminho = gerar_xlsx()
    recalcular_no_excel(caminho)
    print('XLSX:', caminho)
    if args.json:
        exportar_json(args.json, args.separador)
        print('JSON:', args.json)
