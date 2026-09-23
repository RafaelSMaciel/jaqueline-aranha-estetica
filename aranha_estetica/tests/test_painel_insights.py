"""Regressoes do painel analitico/config (auditoria pre-producao, pacote P6).

Dashboard/overview, Excel, busca de clientes, prontuario, anamnese,
configuracoes, branding, financeiro e cancelamento pelo painel.
"""
import io
import json
import re
from datetime import datetime, time, timedelta
from decimal import Decimal
from unittest import mock

from django.core.cache import cache
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from aranha_estetica.models import (
    AceiteTermo,
    AnotacaoSessao,
    Atendimento,
    AvaliacaoNPS,
    Configuracao,
    ConsumoSessao,
    FormularioAnamnese,
    LogAuditoria,
    MovimentoComissao,
    Prontuario,
    ProntuarioVersao,
    RespostaAnamnese,
    Usuario,
    VersaoTermo,
)
from aranha_estetica.utils.branding import get_branding
from aranha_estetica.utils.saude import alertas_saude

from .factories import (
    criar_atendimento,
    criar_cliente,
    criar_compra_pacote,
    criar_pacote,
    criar_procedimento,
    criar_profissional,
)


class _AdminBase(TestCase):
    def setUp(self):
        cache.clear()
        self.admin = Usuario.objects.create_user(
            email='admin@insights.com', password='senha123', papel=Usuario.PAPEL_ADMIN,
        )
        self.client = Client()
        self.client.force_login(self.admin)
        self.prof = criar_profissional()
        self.proc = criar_procedimento(profissional=self.prof)


def _json_script(html: str, elem_id: str):
    m = re.search(rf'<script id="{elem_id}" type="application/json">(.*?)</script>', html, re.S)
    assert m, f'json_script {elem_id} ausente'
    return json.loads(m.group(1))


# ════════════════════════════════════════════════════════════════════
# Dashboard / overview
# ════════════════════════════════════════════════════════════════════
class OverviewTests(_AdminBase):
    def test_grafico_semana_recebe_array_e_nao_string(self):
        resp = self.client.get(reverse('aranha:painel_overview'))
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        dias = _json_script(html, 'id_dias_semana')
        dados = _json_script(html, 'id_dados_grafico_semana')
        self.assertIsInstance(dias, list)
        self.assertEqual(len(dias), 7)
        self.assertIsInstance(dados, list)
        self.assertIsInstance(_json_script(html, 'id_status_labels'), list)

    def test_card_nps_e_nps_real_nao_media(self):
        cli = criar_cliente()
        for nota in (10, 10, 5):
            at = criar_atendimento(cli, self.prof, self.proc, status='REALIZADO')
            AvaliacaoNPS.objects.create(atendimento=at, nota=nota)
        resp = self.client.get(reverse('aranha:painel_overview'))
        # (2 promotores - 1 detrator) / 3 = 33 — a media (8.3) nao e NPS
        self.assertEqual(resp.context['nps_30d'], 33)

    def test_ticket_medio_ignora_retorno_e_sessao_de_pacote(self):
        cli = criar_cliente()
        agora = timezone.now()
        at = criar_atendimento(cli, self.prof, self.proc, data_hora=agora, status='REALIZADO')
        Atendimento.objects.filter(pk=at.pk).update(valor_cobrado=Decimal('200.00'))
        ret = criar_atendimento(cli, self.prof, self.proc, data_hora=agora, status='REALIZADO')
        Atendimento.objects.filter(pk=ret.pk).update(
            valor_cobrado=Decimal('0'), eh_retorno=True, atendimento_origem=at,
        )
        resp = self.client.get(reverse('aranha:painel_overview'))
        self.assertEqual(resp.context['ticket_medio'], '200,00')

    def test_filtro_profissional_nao_numerico_nao_da_500(self):
        resp = self.client.get(reverse('aranha:painel_agendamentos') + '?profissional=abc')
        self.assertEqual(resp.status_code, 200)

    def test_painel_leva_profissional_para_agenda(self):
        user = Usuario.objects.create_user(
            email='prof@insights.com', password='x', papel=Usuario.PAPEL_PROFISSIONAL,
            profissional=self.prof,
        )
        c = Client()
        c.force_login(user)
        resp = c.get(reverse('aranha:painel'))
        self.assertRedirects(resp, reverse('aranha:profissional_agenda'), fetch_redirect_response=False)


class ExportarExcelTests(_AdminBase):
    def _planilha(self):
        import openpyxl
        resp = self.client.get(reverse('aranha:exportar_relatorio_excel'))
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content)).active

    def test_nome_com_formula_vira_texto(self):
        cli = criar_cliente(nome='=HYPERLINK("http://evil.example","clique")')
        ontem = timezone.now() - timedelta(days=1)
        criar_atendimento(cli, self.prof, self.proc, data_hora=ontem, status='REALIZADO')
        ws = self._planilha()
        celula = ws.cell(row=2, column=4)
        self.assertEqual(celula.data_type, 's')
        self.assertTrue(celula.value.startswith('=HYPERLINK'))

    def test_hora_exportada_no_fuso_local(self):
        cli = criar_cliente()
        ontem = timezone.localdate() - timedelta(days=1)
        local_11h = timezone.make_aware(datetime.combine(ontem, time(11, 0)))
        criar_atendimento(cli, self.prof, self.proc, data_hora=local_11h, status='AGENDADO')
        ws = self._planilha()
        self.assertEqual(ws.cell(row=2, column=3).value, '11:00')
        self.assertEqual(ws.cell(row=2, column=7).value, 'Agendado')


# ════════════════════════════════════════════════════════════════════
# Busca de clientes (telefone/CPF com mascara)
# ════════════════════════════════════════════════════════════════════
class BuscaClienteTests(_AdminBase):
    def setUp(self):
        super().setUp()
        self.maria = criar_cliente(nome='Maria Teste', telefone='17999990001', cpf='52998224725')
        self.joana = criar_cliente(nome='Joana Outra', telefone='17988882222')

    def _nomes(self, url, termo):
        resp = self.client.get(url, {'search': termo})
        self.assertEqual(resp.status_code, 200)
        return resp.content.decode()

    def test_clientes_por_telefone_mascarado(self):
        html = self._nomes(reverse('aranha:painel_clientes'), '(17) 99999-0001')
        self.assertIn('Maria Teste', html)
        self.assertNotIn('Joana Outra', html)

    def test_clientes_por_cpf_mascarado(self):
        html = self._nomes(reverse('aranha:painel_clientes'), '529.982.247-25')
        self.assertIn('Maria Teste', html)

    def test_nome_com_digito_nao_vira_busca_por_telefone(self):
        html = self._nomes(reverse('aranha:painel_clientes'), 'Maria 2')
        self.assertNotIn('Joana Outra', html)

    def test_prontuario_busca_nao_da_500_e_acha_mascarado(self):
        # NoReverseMatch 'admin_prontuario' derrubava qualquer busca
        html = self._nomes(reverse('aranha:prontuario_consentimento'), '(17) 99999-0001')
        self.assertIn('Maria Teste', html)
        html = self._nomes(reverse('aranha:prontuario_consentimento'), 'zzz-inexistente')
        self.assertIn('Limpar busca', html)


# ════════════════════════════════════════════════════════════════════
# Prontuario
# ════════════════════════════════════════════════════════════════════
class ProntuarioTermosTests(_AdminBase):
    def test_termos_separados_por_tipo(self):
        cli = criar_cliente()
        hoje = timezone.localdate()
        lgpd = VersaoTermo.objects.create(
            tipo='LGPD', titulo='Privacidade', conteudo='x', versao='1', vigente_desde=hoje,
        )
        termo_proc = VersaoTermo.objects.create(
            tipo='PROCEDIMENTO', procedimento=self.proc, titulo='Termo Limpeza',
            conteudo='y', versao='2', vigente_desde=hoje,
        )
        AceiteTermo.objects.create(cliente=cli, versao_termo=lgpd)
        AceiteTermo.objects.create(cliente=cli, versao_termo=termo_proc)
        resp = self.client.get(reverse('aranha:prontuario_detalhe', args=[cli.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual([a.versao_termo_id for a in resp.context['aceites']], [lgpd.pk])
        self.assertEqual([a.versao_termo_id for a in resp.context['assinaturas']], [termo_proc.pk])


# ════════════════════════════════════════════════════════════════════
# Anamnese (editor)
# ════════════════════════════════════════════════════════════════════
class AnamneseEditorTests(_AdminBase):
    def setUp(self):
        super().setUp()
        self.form = FormularioAnamnese.objects.create(
            nome='Pesquisa pós', tipo='PESQUISA', escopo='GLOBAL',
            schema_json=[{'key': 'nota', 'tipo': 'text', 'label': 'Como foi?', 'obrigatorio': True}],
        )
        self.url = reverse('aranha:admin_anamnese_editar', args=[self.form.pk])

    def test_textarea_e_json_valido_e_reenvio_preserva_pesquisa(self):
        resp = self.client.get(self.url)
        texto = resp.context['schema_texto']
        self.assertEqual(json.loads(texto), self.form.schema_json)
        resp = self.client.post(self.url, {
            'nome': 'Pesquisa pós', 'tipo': 'PESQUISA', 'escopo': 'GLOBAL',
            'schema_json': texto, 'ativo': '1',
        })
        self.assertEqual(resp.status_code, 302)
        self.form.refresh_from_db()
        self.assertEqual(self.form.tipo, 'PESQUISA')

    def test_post_sem_tipo_nao_rebaixa_para_anamnese(self):
        self.client.post(self.url, {
            'nome': 'Pesquisa pós', 'escopo': 'GLOBAL',
            'schema_json': json.dumps(self.form.schema_json), 'ativo': '1',
        })
        self.form.refresh_from_db()
        self.assertEqual(self.form.tipo, 'PESQUISA')

    def test_schema_invalido_e_rejeitado(self):
        resp = self.client.post(self.url, {
            'nome': 'Pesquisa pós', 'tipo': 'PESQUISA', 'escopo': 'GLOBAL',
            'schema_json': '[{"foo": 1}, "x"]', 'ativo': '1',
        })
        self.assertEqual(resp.status_code, 200)  # re-render com erro
        self.assertEqual(resp.context['schema_texto'], '[{"foo": 1}, "x"]')
        self.form.refresh_from_db()
        self.assertEqual(self.form.schema_json[0]['key'], 'nota')

    def test_escopo_modalidade_grava_modalidade(self):
        self.client.post(self.url, {
            'nome': 'Online', 'tipo': 'ANAMNESE', 'escopo': 'MODALIDADE', 'modalidade': 'ONLINE',
            'schema_json': json.dumps(self.form.schema_json), 'ativo': '1',
        })
        self.form.refresh_from_db()
        self.assertEqual((self.form.escopo, self.form.modalidade), ('MODALIDADE', 'ONLINE'))

    def test_excluir_com_respostas_desativa_em_vez_de_500(self):
        RespostaAnamnese.objects.create(formulario=self.form, cliente=criar_cliente(), respostas_json={})
        resp = self.client.post(reverse('aranha:admin_anamnese_excluir', args=[self.form.pk]))
        self.assertEqual(resp.status_code, 302)
        self.form.refresh_from_db()
        self.assertFalse(self.form.ativo)

    def test_criar_com_exemplo_padrao(self):
        url = reverse('aranha:admin_anamnese_criar')
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        exemplo = json.loads(resp.context['schema_texto'])
        self.assertIn('obrigatorio', exemplo[0])  # sem acento: e o que o renderer publico le
        resp = self.client.post(url, {
            'nome': 'Nova', 'tipo': 'ANAMNESE', 'escopo': 'GLOBAL',
            'schema_json': resp.context['schema_texto'], 'ativo': '1',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(FormularioAnamnese.objects.filter(nome='Nova', tipo='ANAMNESE').exists())

    def test_excluir_sem_respostas_apaga(self):
        self.client.post(reverse('aranha:admin_anamnese_excluir', args=[self.form.pk]))
        self.assertFalse(FormularioAnamnese.objects.filter(pk=self.form.pk).exists())


# ════════════════════════════════════════════════════════════════════
# Configuracoes / Branding
# ════════════════════════════════════════════════════════════════════
class ConfiguracoesTests(_AdminBase):
    def test_chave_minuscula_e_preservada(self):
        self.client.post(reverse('aranha:admin_criar_configuracao'), {
            'chave': 'email_admin', 'valor': 'dona@clinica.com',
        })
        self.assertTrue(Configuracao.objects.filter(chave='email_admin').exists())
        self.assertFalse(Configuracao.objects.filter(chave='EMAIL_ADMIN').exists())

    def test_perguntas_prontuario_json_invalido_rejeitado(self):
        self.client.post(reverse('aranha:admin_criar_configuracao'), {
            'chave': 'prontuario_perguntas', 'valor': '{nao e json',
        })
        self.assertFalse(Configuracao.objects.filter(chave='prontuario_perguntas').exists())

    def test_sugestoes_so_chaves_lidas_pelo_codigo(self):
        resp = self.client.get(reverse('aranha:admin_configuracoes'))
        chaves = {s['chave'] for s in resp.context['sugestoes']}
        self.assertEqual(chaves, {'email_admin', 'prontuario_perguntas'})


class BrandingTests(_AdminBase):
    def _post(self, **dados):
        return self.client.post(reverse('aranha:admin_branding'), dados)

    def test_salvar_telefone_chega_ao_site(self):
        cor_atual = get_branding()['THEME_COLOR']
        self._post(CLINIC_PHONE='(17) 3333-4444', THEME_COLOR=cor_atual)
        self.assertEqual(get_branding()['CLINIC_PHONE'], '(17) 3333-4444')
        # cor igual a em uso nao vira override silencioso
        self.assertFalse(Configuracao.objects.filter(chave='THEME_COLOR').exists())
        home = self.client.get(reverse('aranha:inicio'))
        self.assertEqual(home.context['CLINIC_PHONE'], '(17) 3333-4444')

    def test_whatsapp_invalido_nao_grava(self):
        self._post(WHATSAPP_NUMERO='123')
        self.assertFalse(Configuracao.objects.filter(chave='WHATSAPP_NUMERO').exists())

    def test_sem_upload_de_logo(self):
        resp = self.client.get(reverse('aranha:admin_branding'))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'name="logo"')


# ════════════════════════════════════════════════════════════════════
# Financeiro: sessao de pacote nao conta a preco cheio; venda de pacote conta
# ════════════════════════════════════════════════════════════════════
class FinanceiroPacoteTests(_AdminBase):
    def test_receita_de_pacote_na_venda_e_nao_na_sessao(self):
        cli = criar_cliente()
        agora = timezone.now()
        # 2o profissional: no Postgres o EXCLUDE excl_atendimento_sobreposicao
        # recusa 2 AGENDADO do mesmo profissional no mesmo horario.
        outra = criar_profissional(nome='Dra. Bia')
        avulso = criar_atendimento(cli, self.prof, self.proc, data_hora=agora, status='AGENDADO')
        sessao = criar_atendimento(cli, outra, self.proc, data_hora=agora, status='AGENDADO')
        pacote = criar_pacote(preco=Decimal('700.00'), procedimento=self.proc, sessoes=5)
        compra = criar_compra_pacote(cli, pacote)
        ConsumoSessao.objects.create(compra_pacote=compra, atendimento=sessao)
        Atendimento.objects.filter(pk__in=[avulso.pk, sessao.pk]).update(
            status='REALIZADO', valor_cobrado=Decimal('200.00'),
        )
        resp = self.client.get(reverse('aranha:dashboard_financeiro'))
        self.assertEqual(resp.status_code, 200)
        # 200 (avulso) + 700 (venda do pacote); a sessao (200) nao soma de novo
        self.assertEqual(resp.context['fat_hoje_total'], Decimal('900.00'))
        self.assertEqual(resp.context['fat_hoje_count'], 1)
        self.assertEqual(resp.context['fat_hoje_pacotes'], 1)


# ════════════════════════════════════════════════════════════════════
# Cancelamento pelo painel avisa o cliente
# ════════════════════════════════════════════════════════════════════
class CancelamentoPainelTests(_AdminBase):
    def test_cancelar_via_status_enfileira_email_com_hora_local(self):
        cli = criar_cliente(email='cli@x.com')
        amanha = timezone.localdate() + timedelta(days=1)
        local_14h = timezone.make_aware(datetime.combine(amanha, time(14, 0)))
        at = criar_atendimento(cli, self.prof, self.proc, data_hora=local_14h, status='AGENDADO')
        with mock.patch('aranha_estetica.tasks.send_email_async.delay') as delay, \
                self.captureOnCommitCallbacks(execute=True):
            resp = self.client.post(
                reverse('aranha:admin_atualizar_status'),
                data=json.dumps({'atendimento_id': at.pk, 'status': 'CANCELADO'}),
                content_type='application/json',
            )
        self.assertEqual(resp.status_code, 200, resp.content)
        delay.assert_called_once()
        funcao, email, dados = delay.call_args.args
        self.assertEqual((funcao, email), ('enviar_cancelamento_email', 'cli@x.com'))
        self.assertIn('14:00', dados['data_hora'])


# ════════════════════════════════════════════════════════════════════
# Casca do painel (painel/base_v2.html)
# ════════════════════════════════════════════════════════════════════
class PainelShellTests(_AdminBase):
    def _html(self):
        resp = self.client.get(reverse('aranha:painel_overview'))
        self.assertEqual(resp.status_code, 200)
        return resp.content.decode()

    def test_body_sem_x_cloak_e_backdrop_com(self):
        html = self._html()
        body = re.search(r'<body[^>]*>', html).group(0)
        self.assertNotIn('x-cloak', body)
        self.assertRegex(html, r'x-show="menuAberto"\s+x-cloak')

    def test_busca_global_envia_para_clientes(self):
        html = self._html()
        self.assertRegex(
            html,
            r'<form method="get" action="%s" role="search"' % re.escape(reverse('aranha:painel_clientes')),
        )
        self.assertIn('id="adminGlobalSearch"\n          name="search"', html)

    def test_menu_tem_anamneses_e_logout_por_post(self):
        html = self._html()
        self.assertIn(reverse('aranha:admin_anamneses'), html)
        self.assertNotIn('<a href="%s"' % reverse('aranha:usuario_logout'), html)
        self.assertIn('<form method="post" action="%s"' % reverse('aranha:usuario_logout'), html)

    def test_botao_push_so_com_vapid(self):
        with mock.patch.dict('os.environ', {'WEBPUSH_VAPID_PUBLIC_KEY': ''}):
            self.assertNotIn('id="webpushEnable"', self._html())
        with mock.patch.dict('os.environ', {'WEBPUSH_VAPID_PUBLIC_KEY': 'BPk-teste'}):
            self.assertIn('id="webpushEnable"', self._html())



# ════════════════════════════════════════════════════════════════════
# Dado de saude: ficha da cliente chega a quem atende (gap2-01/-10)
# ════════════════════════════════════════════════════════════════════
def _ficha_alergia(cliente, atendimento=None, texto='Lidocaina e latex'):
    form = FormularioAnamnese.objects.create(
        nome='Ficha facial', tipo='ANAMNESE', escopo='GLOBAL',
        schema_json=[
            {'key': 'alergias', 'tipo': 'text', 'label': 'Possui alergias?', 'obrigatorio': False},
            {'key': 'gestante', 'tipo': 'bool', 'label': 'Está gestante?', 'obrigatorio': True},
        ],
    )
    return RespostaAnamnese.objects.create(
        formulario=form, cliente=cliente, atendimento=atendimento,
        respostas_json={'alergias': texto, 'gestante': True},
        respondida_em=timezone.now(),
    )


class SaudeVisivelTests(_AdminBase):
    def setUp(self):
        super().setUp()
        self.cli = criar_cliente(nome='Carla Lima')
        self.at = criar_atendimento(self.cli, self.prof, self.proc, status='PENDENTE')
        self.resposta = _ficha_alergia(self.cli, self.at)

    def test_prontuario_mostra_alerta_e_ficha_com_label(self):
        resp = self.client.get(reverse('aranha:prontuario_detalhe', args=[self.cli.pk]))
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        self.assertIn('Alerta de saúde', html)
        self.assertRegex(html, r'>\s*Lidocaina e latex')  # texto visivel, nao title=
        self.assertIn('Possui alergias?', html)
        self.assertIn('Fichas de avaliação respondidas', html)

    def test_get_nao_cria_prontuario_vazio(self):
        self.client.get(reverse('aranha:prontuario_detalhe', args=[self.cli.pk]))
        self.assertFalse(Prontuario.objects.filter(cliente=self.cli).exists())

    def test_lista_mostra_ficha_online_e_alerta(self):
        vazio = criar_cliente(nome='Beatriz Sem Dados')
        Prontuario.objects.create(cliente=vazio)  # registro vazio nao e prontuario
        resp = self.client.get(reverse('aranha:prontuario_consentimento'))
        html = resp.content.decode()
        itens = {i['cliente'].pk: i for i in resp.context['clientes_list']}
        self.assertTrue(itens[self.cli.pk]['tem_ficha'])
        self.assertFalse(itens[vazio.pk]['tem_prontuario'])
        self.assertIn('Ficha online', html)
        self.assertIn('Sem dados de saúde', html)
        self.assertRegex(html, r'Possui alergias\?:</strong>\s*Lidocaina e latex')

    def test_prontuario_preenchido_ganha_selo(self):
        Prontuario.objects.create(cliente=self.cli, alergias='Dipirona')
        resp = self.client.get(reverse('aranha:prontuario_consentimento'))
        item = next(i for i in resp.context['clientes_list'] if i['cliente'].pk == self.cli.pk)
        self.assertTrue(item['tem_prontuario'])

    def test_agendamentos_e_overview_trazem_alerta(self):
        resp = self.client.get(reverse('aranha:painel_agendamentos'))
        self.assertContains(resp, 'Lidocaina e latex')
        self.assertIn('no-store', resp['Cache-Control'])
        resp = self.client.get(reverse('aranha:painel_overview'))
        self.assertContains(resp, 'Lidocaina e latex')  # pendente de aprovacao

    def test_respostas_com_label_sim_nao_e_trilha(self):
        url = reverse('aranha:admin_anamnese_respostas', args=[self.resposta.formulario_id])
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Possui alergias?')
        self.assertContains(resp, 'Está gestante?')
        self.assertNotContains(resp, '>True<')
        self.assertIn('no-store', resp['Cache-Control'])
        log = LogAuditoria.objects.get(tabela='resposta_anamnese')
        self.assertEqual(log.detalhes['clientes'], [self.cli.pk])

    def test_editar_schema_com_respostas_cria_nova_versao(self):
        form = self.resposta.formulario
        novo_schema = [{'key': 'alergia_medicamento', 'tipo': 'text', 'label': 'Alergia a remédio?'}]
        resp = self.client.post(reverse('aranha:admin_anamnese_editar', args=[form.pk]), {
            'nome': form.nome, 'tipo': 'ANAMNESE', 'escopo': 'GLOBAL',
            'schema_json': json.dumps(novo_schema), 'ativo': '1',
        })
        self.assertEqual(resp.status_code, 302)
        form.refresh_from_db()
        self.assertFalse(form.ativo)
        self.assertEqual(form.schema_json[0]['key'], 'alergias')  # antigo intacto
        novo = FormularioAnamnese.objects.exclude(pk=form.pk).get(nome=form.nome)
        self.assertTrue(novo.ativo)
        self.assertEqual(novo.schema_json, novo_schema)


# ════════════════════════════════════════════════════════════════════
# prontuario_salvar: POST parcial nao apaga, trava otimista, trilha (gap2-04/-06/-11)
# ════════════════════════════════════════════════════════════════════
class ProntuarioSalvarTests(_AdminBase):
    def setUp(self):
        super().setUp()
        self.cli = criar_cliente(nome='Carla Lima')
        self.url = reverse('aranha:prontuario_salvar', args=[self.cli.pk])

    def _versao(self):
        return self.client.get(
            reverse('aranha:prontuario_detalhe', args=[self.cli.pk])
        ).context['versao_prontuario']

    def test_post_parcial_nao_apaga_alergias(self):
        Prontuario.objects.create(cliente=self.cli, alergias='Lidocaina', contraindicacoes='Gestante')
        resp = self.client.post(self.url, {'observacoes_gerais': 'Pele sensível', 'versao': self._versao()})
        self.assertEqual(resp.status_code, 302)
        p = Prontuario.objects.get(cliente=self.cli)
        self.assertEqual((p.alergias, p.contraindicacoes, p.observacoes_gerais),
                         ('Lidocaina', 'Gestante', 'Pele sensível'))

    def test_versao_desatualizada_e_recusada(self):
        Prontuario.objects.create(cliente=self.cli, alergias='Lidocaina')
        versao = self._versao()
        self.client.post(self.url, {'alergias': 'Lidocaina, latex', 'versao': versao})
        # 2a pessoa com a pagina antiga aberta
        self.client.post(self.url, {'alergias': '', 'versao': versao})
        self.assertEqual(Prontuario.objects.get(cliente=self.cli).alergias, 'Lidocaina, latex')

    def test_trilha_usa_cliente_pk_sem_nome_e_so_campos(self):
        self.client.get(reverse('aranha:prontuario_detalhe', args=[self.cli.pk]))
        self.client.post(self.url, {'alergias': 'Dipirona', 'versao': ''})
        logs = LogAuditoria.objects.filter(tabela='prontuario')
        self.assertEqual({lg.registro_id for lg in logs}, {self.cli.pk})
        self.assertFalse(any('Carla' in lg.acao for lg in logs))
        upd = logs.get(acao='Atualizou prontuario')
        self.assertEqual(upd.detalhes['campos_alterados'], ['alergias'])
        self.assertNotIn('Dipirona', json.dumps(upd.detalhes))

    def test_anotacao_com_autor_quebra_de_linha_e_paginacao(self):
        base = timezone.now() - timedelta(days=400)
        for i in range(31):
            criar_atendimento(self.cli, self.prof, self.proc,
                              data_hora=base + timedelta(days=i), status='REALIZADO')
        antigo = Atendimento.objects.filter(cliente=self.cli).order_by('data_hora_inicio').first()
        AnotacaoSessao.objects.create(atendimento=antigo, autor=self.admin,
                                      texto='Aplicado 20U glabela.\nSem intercorrencias.')
        detalhe = reverse('aranha:prontuario_detalhe', args=[self.cli.pk])
        self.assertNotContains(self.client.get(detalhe), 'glabela')
        resp = self.client.get(detalhe + '?page=2')
        self.assertContains(resp, 'Aplicado 20U glabela.<br>Sem intercorrencias.')
        self.assertContains(resp, self.admin.email)  # autor (nome vazio -> e-mail)

    def test_anotacao_criada_fica_na_trilha(self):
        at = criar_atendimento(self.cli, self.prof, self.proc, status='AGENDADO')
        resp = self.client.post(reverse('aranha:anotacao_sessao_salvar', args=[at.pk]), {'texto': 'ok'})
        self.assertEqual(resp.status_code, 200)
        anot = AnotacaoSessao.objects.get(atendimento=at)
        self.assertTrue(LogAuditoria.objects.filter(
            tabela='anotacao_sessao', registro_id=anot.pk, acao='Criou anotacao').exists())


# ════════════════════════════════════════════════════════════════════
# Acesso do profissional ao prontuario (gap2-05)
# ════════════════════════════════════════════════════════════════════
class ProntuarioProfissionalTests(_AdminBase):
    def setUp(self):
        super().setUp()
        self.cli = criar_cliente(nome='Dora')
        self.beto = criar_profissional(nome='Beto')
        self.user_beto = Usuario.objects.create_user(
            email='beto@insights.com', password='x', nome='Beto',
            papel=Usuario.PAPEL_PROFISSIONAL, profissional=self.beto,
        )
        self.prof_client = Client()
        self.prof_client.force_login(self.user_beto)
        self.detalhe = reverse('aranha:prontuario_detalhe', args=[self.cli.pk])
        self.salvar = reverse('aranha:prontuario_salvar', args=[self.cli.pk])
        Prontuario.objects.create(cliente=self.cli, alergias='Lidocaina')

    def _at(self, status, dias, prof=None):
        quando = (timezone.now() + timedelta(days=dias)).replace(hour=10, minute=0, second=0, microsecond=0)
        return criar_atendimento(self.cli, prof or self.beto, self.proc, data_hora=quando, status=status)

    def test_so_cancelado_nao_da_acesso(self):
        self._at('CANCELADO', 2)
        self.assertEqual(self.prof_client.get(self.detalhe).status_code, 403)
        self.assertEqual(self.prof_client.post(self.salvar, {'alergias': ''}).status_code, 403)
        self.assertEqual(Prontuario.objects.get(cliente=self.cli).alergias, 'Lidocaina')
        self.assertTrue(LogAuditoria.objects.filter(acao='Acesso NEGADO a prontuario').exists())

    def test_sem_vinculo_nao_da_acesso(self):
        self.assertEqual(self.prof_client.get(self.detalhe).status_code, 403)

    def test_pendente_futuro_le_mas_nao_grava(self):
        self._at('PENDENTE', 2)
        resp = self.prof_client.get(self.detalhe)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Lidocaina')
        self.assertFalse(resp.context['pode_editar'])
        self.assertEqual(self.prof_client.post(self.salvar, {'alergias': ''}).status_code, 403)
        self.assertEqual(Prontuario.objects.get(cliente=self.cli).alergias, 'Lidocaina')

    def test_realizado_recente_le_e_grava(self):
        self._at('REALIZADO', -3)
        self.assertEqual(self.prof_client.get(self.detalhe).status_code, 200)
        resp = self.prof_client.post(self.salvar, {'observacoes_gerais': 'ok'})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Prontuario.objects.get(cliente=self.cli).observacoes_gerais, 'ok')

    def test_realizado_antigo_nao_da_acesso(self):
        self._at('REALIZADO', -400)
        self.assertEqual(self.prof_client.get(self.detalhe).status_code, 403)

    def test_portal_sem_menu_do_painel_e_nota_so_no_proprio(self):
        meu = self._at('AGENDADO', 1)
        de_outra = self._at('REALIZADO', -5, prof=self.prof)
        resp = self.prof_client.get(self.detalhe)
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        self.assertNotIn(reverse('aranha:painel_agendamentos'), html)
        self.assertIn(reverse('aranha:profissional_agenda'), html)
        self.assertIn(f'data-atendimento-id="{meu.pk}"', html)
        self.assertNotIn(f'data-atendimento-id="{de_outra.pk}"', html)
        self.assertIn('no-store', resp['Cache-Control'])


# ════════════════════════════════════════════════════════════════════
# Diversos do pacote: baixa de comissao na trilha, WhatsApp com DDI, porta do painel
# ════════════════════════════════════════════════════════════════════
class PainelDiversosTests(_AdminBase):
    def test_baixa_de_comissao_registra_auditoria(self):
        at = criar_atendimento(criar_cliente(), self.prof, self.proc, status='AGENDADO')
        mov = MovimentoComissao.objects.create(
            profissional=self.prof, atendimento=at, valor=Decimal('10.00'),
            status=MovimentoComissao.STATUS_PENDENTE,
        )
        self.client.post(reverse('aranha:admin_comissao_pagar', args=[mov.pk]))
        self.assertTrue(LogAuditoria.objects.filter(
            tabela='movimento_comissao', registro_id=mov.pk).exists())

    def test_whatsapp_sem_ddi_ganha_55(self):
        self.client.post(reverse('aranha:admin_branding'), {'WHATSAPP_NUMERO': '(17) 99123-4567'})
        self.assertEqual(Configuracao.objects.get(chave='WHATSAPP_NUMERO').valor, '5517991234567')
        self.assertTrue(LogAuditoria.objects.filter(tabela='configuracao').exists())

    def test_nao_staff_com_profissional_ativo_vai_ao_portal(self):
        user = Usuario.objects.create_user(
            email='recep@insights.com', password='x', nome='R',
            papel=Usuario.PAPEL_RECEPCAO, profissional=self.prof,
        )
        c = Client()
        c.force_login(user)
        resp = c.get(reverse('aranha:painel'))
        self.assertRedirects(resp, reverse('aranha:profissional_agenda'), fetch_redirect_response=False)


# ════════════════════════════════════════════════════════════════════
# Wave 3 — atualizar-status: termo pendente e entrada invalida
# ════════════════════════════════════════════════════════════════════
class AtualizarStatusTermoTests(_AdminBase):
    """REALIZADO com termo do procedimento pendente exige override auditado (rev_painel-01/crawl-2)."""

    def setUp(self):
        super().setUp()
        self.cli = criar_cliente(nome='Rita')
        self.at = criar_atendimento(self.cli, self.prof, self.proc, status='AGENDADO')
        self.termo = VersaoTermo.objects.create(
            tipo='PROCEDIMENTO', procedimento=self.proc, titulo='Termo do procedimento',
            conteudo='Riscos e cuidados.', versao='1.0', vigente_desde=timezone.localdate(),
        )
        self.url = reverse('aranha:admin_atualizar_status')

    def _post(self, **extra):
        payload = {'atendimento_id': self.at.pk, 'status': 'REALIZADO', **extra}
        return self.client.post(self.url, data=json.dumps(payload), content_type='application/json')

    def test_sem_override_recusa_e_nao_muda_status(self):
        resp = self._post()
        self.assertEqual(resp.status_code, 409)
        self.assertTrue(resp.json()['termo_pendente'])
        self.at.refresh_from_db()
        self.assertEqual(self.at.status, 'AGENDADO')
        self.assertFalse(LogAuditoria.objects.filter(acao='Realizado sem termo aceito').exists())

    def test_override_so_com_booleano_true(self):
        self.assertEqual(self._post(sem_termo='true').status_code, 409)
        self.assertEqual(self._post(sem_termo=1).status_code, 409)

    def test_override_explicito_marca_e_audita(self):
        resp = self._post(sem_termo=True)
        self.assertEqual(resp.status_code, 200, resp.content)
        self.at.refresh_from_db()
        self.assertEqual(self.at.status, 'REALIZADO')
        log = LogAuditoria.objects.get(acao='Realizado sem termo aceito')
        self.assertEqual((log.tabela, log.registro_id), ('atendimento', self.at.pk))
        self.assertEqual(log.detalhes['termos'], [self.termo.pk])
        self.assertEqual(log.usuario_id, self.admin.pk)
        self.assertIsNotNone(log.ip_origem)

    def test_termo_aceito_nao_pede_override(self):
        AceiteTermo.objects.create(cliente=self.cli, versao_termo=self.termo, atendimento=self.at)
        resp = self._post()
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertFalse(LogAuditoria.objects.filter(acao='Realizado sem termo aceito').exists())

    def test_outros_status_nao_dependem_do_termo(self):
        resp = self.client.post(self.url, data=json.dumps(
            {'atendimento_id': self.at.pk, 'status': 'CONFIRMADO'}), content_type='application/json')
        self.assertEqual(resp.status_code, 200)

    def test_overview_trata_termo_pendente(self):
        html = self.client.get(reverse('aranha:painel_overview')).content.decode()
        self.assertIn('data.termo_pendente', html)
        self.assertIn('payload.sem_termo = true', html)


class AtualizarStatusEntradaInvalidaTests(_AdminBase):
    """Entrada invalida -> 400/404 sem ERROR/traceback (crawl-3)."""

    def test_payloads_invalidos_nao_viram_500(self):
        at = criar_atendimento(criar_cliente(), self.prof, self.proc, status='AGENDADO')
        url = reverse('aranha:admin_atualizar_status')
        casos = [
            (json.dumps({'atendimento_id': None, 'status': 'CONFIRMADO'}), 400),
            (json.dumps({'atendimento_id': 999999, 'status': 'CONFIRMADO'}), 404),
            (json.dumps({'atendimento_id': 'abc', 'status': 'CONFIRMADO'}), 400),
            (json.dumps({'atendimento_id': True, 'status': 'CONFIRMADO'}), 400),
            (json.dumps(['x']), 400),
            (json.dumps({'atendimento_id': at.pk, 'status': None}), 400),
            (json.dumps({'atendimento_id': at.pk, 'status': 5}), 400),
            (json.dumps({'atendimento_id': at.pk}), 400),
            ('{nao e json', 400),
            (b'\xff\xfe', 400),
        ]
        with self.assertNoLogs('aranha_estetica.views.admin', level='ERROR'):
            for corpo, esperado in casos:
                with self.subTest(corpo=corpo):
                    resp = self.client.post(url, data=corpo, content_type='application/json')
                    self.assertEqual(resp.status_code, esperado, resp.content)
                    self.assertIn('erro', resp.json())
        at.refresh_from_db()
        self.assertEqual(at.status, 'AGENDADO')

    def test_id_como_texto_numerico_continua_valendo(self):
        at = criar_atendimento(criar_cliente(), self.prof, self.proc, status='AGENDADO')
        resp = self.client.post(
            reverse('aranha:admin_atualizar_status'),
            data=json.dumps({'atendimento_id': str(at.pk), 'status': ' confirmado '}),
            content_type='application/json',
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        at.refresh_from_db()
        self.assertEqual(at.status, 'CONFIRMADO')


# ════════════════════════════════════════════════════════════════════
# Wave 3 — no-show do financeiro conta todos os realizados (rev_painel-04)
# ════════════════════════════════════════════════════════════════════
class FinanceiroNoShowTests(_AdminBase):
    def _realizado(self, cli, dias, compra=None, **campos):
        # AGENDADO + update (sem signals): o sinal de REALIZADO consumiria o pacote sozinho
        quando = timezone.now() - timedelta(days=dias)
        at = criar_atendimento(cli, self.prof, self.proc, data_hora=quando, status='AGENDADO')
        if compra is not None:
            ConsumoSessao.objects.create(compra_pacote=compra, atendimento=at)
        Atendimento.objects.filter(pk=at.pk).update(status='REALIZADO', **campos)
        return at

    def test_sessoes_de_pacote_entram_no_denominador(self):
        cli = criar_cliente()
        compra = criar_compra_pacote(cli, criar_pacote(procedimento=self.proc, sessoes=10))
        for dia in range(1, 9):
            self._realizado(cli, dia, compra=compra, valor_cobrado=Decimal('200.00'))
        self._realizado(cli, 10, valor_cobrado=Decimal('200.00'))  # avulso
        criar_atendimento(cli, self.prof, self.proc,
                          data_hora=timezone.now() - timedelta(days=11), status='FALTOU')
        ctx = self.client.get(reverse('aranha:dashboard_financeiro')).context
        self.assertEqual(ctx['fat_mes_count'], 1)  # faturamento continua so o avulso
        self.assertEqual(ctx['no_show_count'], 1)
        self.assertEqual(ctx['no_show_pct'], 10.0)

    def test_retorno_e_sem_valor_entram_no_denominador(self):
        cli = criar_cliente()
        origem = self._realizado(cli, 20, valor_cobrado=Decimal('150.00'))
        self._realizado(cli, 1, eh_retorno=True, atendimento_origem=origem, valor_cobrado=Decimal('0'))
        self._realizado(cli, 2, valor_cobrado=None)
        self._realizado(cli, 3, valor_cobrado=Decimal('150.00'))
        criar_atendimento(cli, self.prof, self.proc,
                          data_hora=timezone.now() - timedelta(days=4), status='FALTOU')
        ctx = self.client.get(reverse('aranha:dashboard_financeiro')).context
        self.assertEqual(ctx['no_show_pct'], 20.0)  # 1 falta / (4 realizados + 1)


# ════════════════════════════════════════════════════════════════════
# Wave 3 — alertas de saude leem respostas_extras (pgupgrade-02)
# ════════════════════════════════════════════════════════════════════
class AlertasRespostasExtrasTests(_AdminBase):
    def test_respostas_extras_migradas_viram_alerta(self):
        Configuracao.objects.create(chave='prontuario_perguntas', valor=json.dumps([
            {'chave': 'esta_gravida_ou_amamentando', 'texto': 'Está grávida ou amamentando?', 'tipo': 'BOOLEAN'},
        ]))
        cli = criar_cliente(nome='Zélia')
        Prontuario.objects.create(cliente=cli, respostas_extras={
            'esta_gravida_ou_amamentando': True,
            'usa_acido_retinoico': True,
            'fumante': True,
            'doenca_x': False,
            'observacoes_de_pele': 'Sensível',
        })
        alertas = alertas_saude(cli)
        self.assertEqual({a['label'] for a in alertas},
                         {'Está grávida ou amamentando?', 'Usa acido retinoico'})
        self.assertTrue(all(a['origem'] == 'prontuario' and a['valor'] == 'Sim' for a in alertas))
        resp = self.client.get(reverse('aranha:prontuario_detalhe', args=[cli.pk]))
        self.assertContains(resp, 'Alerta de saúde')
        self.assertContains(resp, 'Usa acido retinoico')

    def test_schema_invalido_nao_quebra_alerta(self):
        Configuracao.objects.create(chave='prontuario_perguntas', valor='{nao e json')
        cli = criar_cliente()
        Prontuario.objects.create(cliente=cli, respostas_extras={'esta_gravida_x': True})
        self.assertEqual([a['label'] for a in alertas_saude(cli)], ['Esta gravida x'])


# ════════════════════════════════════════════════════════════════════
# Wave 3 — prontuario: autor snapshot, historico e registro vazio
# ════════════════════════════════════════════════════════════════════
class ProntuarioAutorNotaTests(_AdminBase):
    """Nota com autor excluido nao derruba a tela e mostra o snapshot (followups-anotar-autor-500)."""

    def test_detalhe_usa_autor_nome_e_nao_quebra_sem_autor(self):
        cli = criar_cliente()
        at = criar_atendimento(cli, self.prof, self.proc,
                               data_hora=timezone.now() - timedelta(days=2), status='REALIZADO')
        AnotacaoSessao.objects.create(atendimento=at, autor=None, autor_nome='Dra. Antiga', texto='nota 1')
        AnotacaoSessao.objects.create(atendimento=at, autor=None, autor_nome='', texto='nota 2')
        resp = self.client.get(reverse('aranha:prontuario_detalhe', args=[cli.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Dra. Antiga')
        self.assertContains(resp, 'autor removido', count=1)


class ProntuarioVazioTests(_AdminBase):
    """POST sem alteracao nao deixa Prontuario vazio (tirava a cliente da purga LGPD)."""

    def setUp(self):
        super().setUp()
        self.cli = criar_cliente()
        self.url = reverse('aranha:prontuario_salvar', args=[self.cli.pk])

    def test_post_sem_alteracao_nao_cria_registro(self):
        resp = self.client.post(self.url, {'versao': '', 'alergias': '', 'observacoes_gerais': '  '})
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(Prontuario.objects.filter(cliente=self.cli).exists())

    def test_versao_recusada_nao_cria_registro(self):
        resp = self.client.post(self.url, {'versao': '2020-01-01T00:00:00', 'alergias': 'Dipirona'})
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(Prontuario.objects.filter(cliente=self.cli).exists())


class ProntuarioHistoricoTests(_AdminBase):
    """Edicao guarda o valor anterior em ProntuarioVersao e o detalhe mostra (followups-prontuario-sem-historico)."""

    def setUp(self):
        super().setUp()
        self.cli = criar_cliente()
        self.detalhe = reverse('aranha:prontuario_detalhe', args=[self.cli.pk])
        self.url = reverse('aranha:prontuario_salvar', args=[self.cli.pk])

    def _salvar(self, **campos):
        versao = self.client.get(self.detalhe).context['versao_prontuario']
        return self.client.post(self.url, {'versao': versao, **campos})

    def test_segunda_edicao_guarda_alergia_anterior(self):
        self._salvar(alergias='Dipirona')
        self.assertEqual(ProntuarioVersao.objects.count(), 0)  # 1o preenchimento: nada a guardar
        self._salvar(alergias='Latex')
        versao = ProntuarioVersao.objects.get()
        self.assertEqual(versao.dados['alergias'], 'Dipirona')
        self.assertEqual(versao.autor_id, self.admin.pk)
        self.assertEqual(Prontuario.objects.get(cliente=self.cli).alergias, 'Latex')

    def test_post_sem_alteracao_nao_gera_versao(self):
        self._salvar(alergias='Dipirona')
        self._salvar(alergias='Dipirona')
        self.assertEqual(ProntuarioVersao.objects.count(), 0)

    def test_detalhe_mostra_antes_e_depois_com_autor(self):
        self.admin.nome = 'Jaqueline'
        self.admin.save(update_fields=['nome'])
        self._salvar(alergias='Dipirona')
        self._salvar(alergias='', medicamentos_uso='Roacutan')
        resp = self.client.get(self.detalhe)
        hist = resp.context['historico_prontuario']
        self.assertEqual(len(hist), 1)
        self.assertEqual(hist[0]['autor'], 'Jaqueline')
        mudancas = {m['campo']: (m['antes'], m['depois']) for m in hist[0]['mudancas']}
        self.assertEqual(mudancas, {
            'Alergias': ('Dipirona', ''),
            'Medicamentos em uso': ('', 'Roacutan'),
        })
        self.assertContains(resp, 'Histórico de alterações (1)')
        self.assertContains(resp, '<del class="text-texto-suave">Dipirona</del>')


class AuditoriaAutorSnapshotTests(_AdminBase):
    """Trilha mostra o snapshot usuario_nome: autor excluido nao vira 'Sistema'."""

    def test_autor_excluido_mostra_snapshot(self):
        LogAuditoria.objects.create(usuario=None, usuario_nome='Dra. Ex <ex@x.com>', acao='Acao antiga')
        LogAuditoria.objects.create(usuario=None, acao='Job noturno')
        resp = self.client.get(reverse('aranha:admin_auditoria'))
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        self.assertIn('Dra. Ex &lt;ex@x.com&gt;', html)
        self.assertRegex(html, r'>Sistema</span>')
