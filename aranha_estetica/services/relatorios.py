"""Relatorio gerencial em Excel — mesmos numeros do painel (faturamento D20, services/financeiro).

Antes a planilha somava toda linha com valor (sessao de pacote a preco cheio, retorno, cancelado
com valor): a soma da coluna nao batia com o card "Faturamento" — justo a conferencia que a dona
(ou a banca) faz. Aqui cada atendimento diz se conta no faturamento e quanto; os pacotes vendidos
entram na aba propria e o resumo fecha com o painel.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

from ..models import Atendimento, CompraPacote
from ..utils.datas import dias_locais, fmt_local
from . import financeiro

COLUNAS_ATENDIMENTOS = ['ID', 'Data', 'Hora', 'Cliente', 'Profissional', 'Procedimento', 'Status',
                        'Valor cobrado (R$)', 'Sessão de pacote', 'Retorno', 'Conta no faturamento',
                        'Valor faturado (R$)']
COLUNAS_PACOTES = ['ID', 'Data da venda', 'Cliente', 'Pacote', 'Status', 'Valor pago (R$)',
                   'Reembolsado (R$)', 'Líquido (R$)']


def _texto_seguro(ws) -> None:
    """Impede formula injection na ultima linha: o openpyxl grava como formula toda str iniciada por
    '=' e o nome da cliente vem do agendamento publico (texto livre)."""
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
            cell.data_type = 's'


def _sim(valor: bool) -> str:
    return 'Sim' if valor else 'Não'


def planilha_gerencial(inicio: date, fim: date):
    """Workbook (openpyxl) com Resumo, Atendimentos e Pacotes vendidos do periodo [inicio, fim]."""
    import openpyxl
    from openpyxl.styles import Font

    periodo_atend = dias_locais('data_hora_inicio', inicio, fim)
    atendimentos = (Atendimento.objects.filter(**periodo_atend)
                    .select_related('cliente', 'profissional', 'procedimento', 'sessao_pacote_vinculada')
                    .order_by('data_hora_inicio'))
    faturados = set(financeiro.atendimentos_faturados(atendimentos).values_list('pk', flat=True))
    compras = (CompraPacote.objects.filter(**dias_locais('criado_em', inicio, fim))
               .select_related('cliente').order_by('criado_em'))
    avulsos = financeiro.resumo_avulsos(Atendimento.objects.filter(**periodo_atend))
    pacotes = financeiro.resumo_pacotes(CompraPacote.objects.filter(**dias_locais('criado_em', inicio, fim)))

    wb = openpyxl.Workbook()
    resumo = wb.active
    resumo.title = 'Resumo'
    linhas = [
        ('Período', f'{inicio:%d/%m/%Y} a {fim:%d/%m/%Y}'),
        ('Atendimentos avulsos faturados', avulsos['qtd']),
        ('Receita de atendimentos avulsos (R$)', float(avulsos['total'])),
        ('Ticket médio dos avulsos (R$)', float(avulsos['ticket_medio'])),
        ('Pacotes vendidos (líquido > 0)', pacotes['qtd']),
        ('Receita de pacotes, líquida de reembolso (R$)', float(pacotes['total'])),
        ('Faturamento do período (R$)', float(avulsos['total'] + pacotes['total'])),
    ]
    for rotulo, valor in linhas:
        resumo.append([rotulo, valor])
    resumo.append([])
    resumo.append(['Faturamento = atendimentos avulsos realizados (sem retorno, sem sessão de pacote, valor > 0) '
                   '+ pacotes vendidos no período, líquidos de reembolso — a mesma regra do painel.'])
    resumo.column_dimensions['A'].width = 48
    resumo.column_dimensions['B'].width = 26

    ws = wb.create_sheet('Atendimentos')
    ws.append(COLUNAS_ATENDIMENTOS)
    total_faturado = Decimal('0')
    for at in atendimentos:
        conta = at.pk in faturados
        valor_faturado = (at.valor_cobrado or Decimal('0')) if conta else Decimal('0')
        total_faturado += valor_faturado
        ws.append([
            at.pk,
            fmt_local(at.data_hora_inicio, '%d/%m/%Y'),
            fmt_local(at.data_hora_inicio, '%H:%M'),
            at.cliente.nome,
            at.profissional.nome,
            at.procedimento.nome,
            at.get_status_display(),
            float(at.valor_cobrado or 0),
            _sim(hasattr(at, 'sessao_pacote_vinculada') and at.sessao_pacote_vinculada is not None),
            _sim(at.eh_retorno),
            _sim(conta),
            float(valor_faturado),
        ])
        _texto_seguro(ws)
    ws.append(['Total faturado', *[''] * (len(COLUNAS_ATENDIMENTOS) - 2), float(total_faturado)])

    wp = wb.create_sheet('Pacotes vendidos')
    wp.append(COLUNAS_PACOTES)
    for compra in compras:
        wp.append([
            compra.pk,
            fmt_local(compra.criado_em, '%d/%m/%Y'),
            compra.cliente.nome,
            compra.pacote_nome,
            compra.get_status_display(),
            float(compra.valor_pago),
            float(compra.valor_reembolsado),
            float(compra.valor_pago - compra.valor_reembolsado),
        ])
        _texto_seguro(wp)
    wp.append(['Total líquido', '', '', '', '', '', '', float(pacotes['total'])])

    for planilha in (ws, wp, resumo):
        for cell in planilha[1]:
            cell.font = Font(bold=True)
    return wb, {'atendimentos': atendimentos.count(), 'pacotes': compras.count()}
